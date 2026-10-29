"""Step 12 · Show and save the predicted paper and its mark scheme.

Saves predicted_paper.docx and predicted_mark_scheme.docx (ready to print) plus .md copies in the reports folder.
Each graph is drawn from its formula y = f(x), so the curve is smooth and exact.
Any problem the checks could not fix is printed in the mark scheme as CHECK.
Solve the paper yourself before trusting it: a human check is the final step.
"""
import matplotlib.pyplot as plt
import numpy as np

from .analysis import BLUE
from .display import show_image, show_markdown
from .generation.numbers import run

NOTE = "Practice paper written by an AI model from past-paper patterns. It is not an official Cambridge paper."


def graph_png(q, qn, reports_dir):
    """Draw the question's graph from its formula y = f(x)."""
    path = reports_dir / f"figure_{q}.png"
    xs = np.linspace(qn.graph_x_from, qn.graph_x_to, 300)
    fig, ax = plt.subplots(figsize=(5, 3.2))
    ax.plot(xs, [run(qn.graph_y_formula, {"x": x}) for x in xs], color=BLUE, lw=2)
    ax.set(title=f"Figure {q}.2", xlabel=qn.graph_x_label, ylabel=qn.graph_y_label)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def save_paper(generated, reports_dir, target_exam):
    """Write the Word and Markdown files. Questions are numbered 1, 2, 3 ... even if one was skipped.
    Returns the total marks of the paper."""
    from docx import Document
    from docx.shared import Inches

    title = f"Predicted Paper 4 - {target_exam}"
    qp, ms = Document(), Document()
    qp.add_heading(title, 0)
    qp.add_paragraph(NOTE)
    ms.add_heading(f"Mark scheme - {title}", 0)
    qp_md, ms_md = [f"# {title}\n*{NOTE}*\n"], [f"# Mark scheme - {title}\n"]

    for q, g in enumerate(generated, 1):
        qn = g["question"]
        para = qp.add_paragraph()
        para.add_run(f"{q}   ").bold = True
        para.add_run(qn.stem)
        qp_md.append(f"\n**{q}** {qn.stem}\n")
        if qn.figure:
            qp.add_paragraph().add_run(f"[Figure {q}.1: {qn.figure}]").italic = True
            qp_md.append(f"*[Figure {q}.1: {qn.figure}]*\n")
        if qn.graph_y_formula:
            path = graph_png(q, qn, reports_dir)
            qp.add_picture(str(path), width=Inches(4.5))
            qp_md.append(f"![Figure {q}.2]({path.name})\n")
        for p in qn.parts:
            qp.add_paragraph(f"{p.label} {p.text}    [{p.marks}]")
            qp_md.append(f"{p.label} {p.text} **[{p.marks}]**\n")
        total = sum(p.marks for p in qn.parts)
        qp.add_paragraph(f"[Total: {total}]")
        qp_md.append(f"*[Total: {total}]*\n")

        table = ms.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, head in zip(table.rows[0].cells, ["Question", "Answer", "Marks"]):
            cell.text = head
        ms_md.append(f"\n| Question {q} | Answer | Marks |\n|---|---|---|")
        for p in qn.parts:
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = f"{q}{p.label}", "\n".join(p.mark_scheme), str(p.marks)
            ms_md.append(f"| {q}{p.label} | {'<br>'.join(p.mark_scheme)} | {p.marks} |")
        if g["problems"]:
            ms.add_paragraph("CHECK: " + "; ".join(g["problems"]))
            ms_md.append("\n**CHECK:** " + "; ".join(g["problems"]) + "\n")

    qp.save(reports_dir / "predicted_paper.docx")
    ms.save(reports_dir / "predicted_mark_scheme.docx")
    (reports_dir / "predicted_paper.md").write_text("\n".join(qp_md), encoding="utf-8")
    (reports_dir / "predicted_mark_scheme.md").write_text("\n".join(ms_md), encoding="utf-8")

    paper_total = sum(p.marks for g in generated for p in g["question"].parts)
    print(f"Total marks: {paper_total}" + ("" if paper_total == 100 else " - not 100, run Step 11 again to retry the missing questions"))
    print("Saved in", reports_dir)
    return paper_total


def show_paper(generated, reports_dir, target_exam):
    """Show the paper (with its graphs) in the notebook, or print it in a terminal."""
    show_markdown(f"# Predicted Paper 4 - {target_exam}\n*{NOTE}*")
    for q, g in enumerate(generated, 1):
        qn = g["question"]
        show_markdown(f"**{q}** {qn.stem}" + (f"\n\n*[Figure {q}.1: {qn.figure}]*" if qn.figure else ""))
        if qn.graph_y_formula:
            show_image(reports_dir / f"figure_{q}.png")
        show_markdown("\n\n".join(f"{p.label} {p.text} **[{p.marks}]**" for p in qn.parts))
